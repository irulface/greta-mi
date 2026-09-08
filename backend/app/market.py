"""Market observations, calculations and analyst-managed intelligence."""
import calendar
import csv
import io
import math
import statistics
import zipfile
from datetime import date, datetime
from decimal import Decimal
from typing import Literal

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field, ConfigDict
from sqlalchemy import select

from .db import MarketPoint, Record, User, now, uid
from .security import PERMISSIONS, audit, current_user, db_session, require, visible

router = APIRouter(prefix='/api/v1/market', tags=['Market intelligence'])
LANDSCAPES = ['Leader', 'Challenger', 'Specialist', 'Emerging Supplier', 'New Entrant', 'Local Supplier', 'Global Supplier']
KIND_PERMISSION = {'rfi': 'rfi:read', 'response': 'rfi:read', 'document': 'repository:read',
                   'supplier': 'supplier:read', 'conversation': 'ai', 'analysis': 'ai', 'advanced_analysis': 'ai', 'external_evidence': 'procurement', 'workflow_package': 'procurement', 'autonomous_agent': 'procurement', 'agent_signal': 'procurement', 'agent_alert': 'procurement'}
PRIVATE_KINDS = {'research_workspace', 'saved_view', 'market_brief', 'market_schedule', 'alert_rule', 'market_alert', 'advanced_analysis', 'external_evidence', 'workflow_package', 'autonomous_agent', 'agent_alert'}


def accessible(user, record):
    if not record or not visible(user, record):
        return False
    if KIND_PERMISSION.get(record.kind, 'market:read') not in PERMISSIONS.get(user.role, set()):
        return False
    if record.kind == 'conversation' and record.owner != user.id:
        return False
    if record.kind == 'response' and user.role == 'Fungsi Pengguna':
        return False
    if record.kind in PRIVATE_KINDS and record.owner != user.id and user.id not in record.data.get('shared_with', []):
        return False
    return True


def get(db, user, id, kind=None, lock=False):
    query = select(Record).where(Record.id == id)
    row = db.scalar(query.with_for_update() if lock else query)
    if not accessible(user, row) or (kind and row.kind != kind):
        raise HTTPException(404, 'Data tidak ditemukan atau tidak dapat diakses.')
    return row


def records(db, user, kind, category='', include_demo=True):
    return [r for r in db.scalars(select(Record).where(Record.kind == kind).order_by(Record.updated_at.desc()))
            if accessible(user, r) and (not category or r.data.get('category_id') == category)
            and (include_demo or not r.data.get('is_demo'))]


def public(row):
    excluded = {'storage_path', 'source_path', 'token_hash', 'path', 'api_key', 'password'}
    return {**{key: value for key, value in row.data.items() if key not in excluded},
            'id': row.id, 'owner': row.owner, 'region': row.region, 'classification': row.classification,
            'updated_at': row.updated_at}


class StrictInput(BaseModel):
    model_config = ConfigDict(extra='forbid', allow_inf_nan=False)


class SeriesInput(StrictInput):
    name: str = Field(min_length=3, max_length=150)
    category_id: str
    commodity: str = Field(min_length=2, max_length=150)
    geography: str = Field(min_length=2, max_length=100)
    currency: str = Field(default='USD', max_length=8)
    unit: str = Field(min_length=1, max_length=50)
    metric_type: Literal['price', 'index', 'lead_time', 'inventory_coverage'] = 'price'
    frequency: Literal['daily', 'weekly', 'monthly', 'quarterly', 'irregular'] = 'monthly'
    source: str = Field(min_length=3, max_length=500)
    notes: str = Field(default='', max_length=5000)


class PointInput(StrictInput):
    date: date
    value: Decimal = Field(ge=-1_000_000_000_000, le=1_000_000_000_000, max_digits=24, decimal_places=8)
    confidence: float = Field(default=1, ge=0, le=1)
    source: str = Field(default='', max_length=500)
    notes: str = Field(default='', max_length=5000)


class PointsInput(StrictInput):
    observations: list[PointInput] = Field(min_length=1, max_length=10000)
    replace_existing: bool = False


def points(db, series, start='', end=''):
    query = select(MarketPoint).where(MarketPoint.series_id == series.id).order_by(MarketPoint.observed_on)
    if start: query = query.where(MarketPoint.observed_on >= start)
    if end: query = query.where(MarketPoint.observed_on <= end)
    rows = [{'date': row.observed_on, 'value': float(row.value), 'confidence': float(row.confidence),
             'source': row.source, 'notes': row.notes, 'updated_at': row.updated_at} for row in db.scalars(query)]
    if not rows and series.data.get('series'):
        rows = [{'date': day, 'value': value, 'confidence': 1, 'source': series.data.get('source', ''), 'notes': ''}
                for day, value in zip(series.data['dates'], series.data['series'])
                if (not start or day >= start) and (not end or day <= end)]
    return rows


def change(current, baseline):
    return round((current - baseline) / abs(baseline) * 100, 4) if baseline else None


def shifted(day, months):
    total = day.year * 12 + day.month - 1 - months
    year, month = total // 12, total % 12 + 1
    return day.replace(year=year, month=month, day=min(day.day, calendar.monthrange(year, month)[1]))


def statistics_for(observations):
    if not observations: return {'count': 0}
    values = [p['value'] for p in observations]
    latest, first = observations[-1], observations[0]
    day = date.fromisoformat(latest['date'])
    result = {'count': len(values), 'latest': latest['value'], 'as_of': latest['date'],
              'high': max(values), 'low': min(values), 'mom': None, 'qoq': None, 'yoy': None, 'cagr': None}
    for label, months in [('mom', 1), ('qoq', 3), ('yoy', 12)]:
        target = shifted(day, months)
        candidates = [p for p in observations if p['date'] <= target.isoformat()]
        baseline = candidates[-1] if candidates else None
        if baseline and (target - date.fromisoformat(baseline['date'])).days <= 40:
            result[label] = change(latest['value'], baseline['value'])
            result[label + '_baseline_date'] = baseline['date']
    elapsed = (day - date.fromisoformat(first['date'])).days
    if elapsed >= 365 and first['value'] > 0 and latest['value'] > 0:
        result['cagr'] = round(((latest['value'] / first['value']) ** (365.25 / elapsed) - 1) * 100, 4)
    for window in [7, 30]:
        result['ma_' + str(window)] = round(statistics.mean(values[-window:]), 6) if len(values) >= window else None
    returns = [change(b, a) for a, b in zip(values, values[1:]) if a]
    result['volatility'] = round(statistics.stdev(returns), 4) if len(returns) >= 2 else None
    result['stale_days'] = max(0, (date.today() - day).days)
    result['average_confidence'] = round(statistics.mean(p['confidence'] for p in observations), 4)
    return result


def series_public(db, series, start='', end='', include_points=False):
    observation_rows = points(db, series, start, end)
    data = public(series)
    data.pop('series', None); data.pop('dates', None)
    data.setdefault('commodity', data['name'])
    data.setdefault('geography', series.data.get('region', series.region))
    data.setdefault('metric_type', 'price'); data.setdefault('frequency', 'monthly')
    data['statistics'] = statistics_for(observation_rows)
    if include_points: data['observations'] = observation_rows
    return data


def upsert_points(db, user, series, body):
    dates = [p.date.isoformat() for p in body.observations]
    if len(set(dates)) != len(dates): raise HTTPException(422, 'Tanggal duplikat dalam satu impor.')
    if any(p.date > date.today() for p in body.observations): raise HTTPException(422, 'Observasi pasar tidak boleh bertanggal di masa depan.')
    if series.data.get('metric_type') in ('lead_time', 'inventory_coverage') and any(p.value < 0 for p in body.observations):
        raise HTTPException(422, 'Lead time dan inventory coverage tidak boleh negatif.')
    existing = {p.observed_on: p for p in db.scalars(select(MarketPoint).where(MarketPoint.series_id == series.id))}
    inserted = updated = unchanged = 0
    conflicts = []
    for p in body.observations:
        old = existing.get(p.date.isoformat())
        values = {'value': p.value, 'confidence': Decimal(str(round(p.confidence, 4))), 'source': p.source or series.data['source'], 'notes': p.notes}
        equal = old and all(getattr(old, key) == value for key, value in values.items())
        if old and not equal and not body.replace_existing:
            conflicts.append(p.date.isoformat())
    if conflicts: raise HTTPException(409, 'Tanggal sudah memiliki nilai berbeda: ' + ', '.join(conflicts[:10]) + '. Pilih replace existing untuk membuat koreksi tercatat.')
    for p in body.observations:
        old = existing.get(p.date.isoformat())
        values = {'value': p.value, 'confidence': Decimal(str(round(p.confidence, 4))), 'source': p.source or series.data['source'], 'notes': p.notes}
        if old and all(getattr(old, key) == value for key, value in values.items()):
            unchanged += 1; continue
        if old:
            audit(db, user, 'market.point.corrected', old.id, {'series_id': series.id, 'date': old.observed_on, 'previous_value': str(old.value), 'new_value': str(p.value)})
            for key, value in values.items(): setattr(old, key, value)
            old.updated_at = now(); updated += 1
        else:
            db.add(MarketPoint(series_id=series.id, observed_on=p.date.isoformat(), **values)); inserted += 1
    db.flush()
    latest = db.scalar(select(MarketPoint).where(MarketPoint.series_id == series.id).order_by(MarketPoint.observed_on.desc()).limit(1))
    series.data = {**series.data, 'value': float(latest.value), 'date': latest.observed_on,
                   'change': statistics_for(points(db, series)).get('mom')}
    series.updated_at = now()
    audit(db, user, 'market.observations.saved', series.id, {'inserted': inserted, 'updated': updated, 'unchanged': unchanged})
    return {'inserted': inserted, 'updated': updated, 'unchanged': unchanged}


@router.get('/series')
def list_series(category: str = '', include_demo: bool = True, user=Depends(current_user), db=Depends(db_session)):
    require(user, 'market:read')
    return [series_public(db, row) for row in records(db, user, 'price', category, include_demo)]


@router.post('/series')
def create_series(body: SeriesInput, user=Depends(current_user), db=Depends(db_session)):
    require(user, 'market:write'); get(db, user, body.category_id, 'category')
    if body.metric_type == 'price' and (len(body.currency) != 3 or not body.currency.isalpha()):
        raise HTTPException(422, 'Gunakan kode mata uang tiga huruf untuk harga.')
    data = body.model_dump(); data['currency'] = body.currency.upper() if body.metric_type == 'price' else ''
    row = Record(kind='price', owner=user.id, region=user.region, data=data)
    db.add(row); db.flush(); audit(db, user, 'market.series.create', row.id); db.commit()
    return series_public(db, row)


@router.get('/series/{id}')
def series_detail(id: str, start: date | None = None, end: date | None = None, user=Depends(current_user), db=Depends(db_session)):
    require(user, 'market:read'); row = get(db, user, id, 'price')
    if start and end and start > end: raise HTTPException(422, 'Rentang tanggal tidak valid.')
    return series_public(db, row, start.isoformat() if start else '', end.isoformat() if end else '', True)


@router.post('/series/{id}/observations')
def add_observations(id: str, body: PointsInput, user=Depends(current_user), db=Depends(db_session)):
    require(user, 'market:write'); row = get(db, user, id, 'price', lock=True)
    if row.data.get('is_demo'): raise HTTPException(409, 'Buat seri nyata untuk menyimpan observasi; seri contoh tidak dapat diubah.')
    result = upsert_points(db, user, row, body); db.commit(); return result


def parse_import(content, filename):
    try:
        if filename.lower().endswith('.csv'):
            reader = csv.DictReader(io.StringIO(content.decode('utf-8-sig')))
            header = [str(value).strip() for value in (reader.fieldnames or [])]
            if len(header) != len(set(header)): raise ValueError('Duplicate header')
            table = list(reader)
            if any(None in row for row in table): raise ValueError('Extra columns without a header')
        elif filename.lower().endswith('.xlsx'):
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 2000 or sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError('Expanded workbook exceeds the import limit')
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                rows = workbook.active.iter_rows(values_only=True)
                header = [str(v).strip() for v in next(rows)]
                if len(set(header)) != len(header): raise ValueError('Duplicate header')
                table = []
                for cells in rows:
                    if any(isinstance(v, str) and v.startswith('=') for v in cells): raise ValueError('Formula cells are not accepted')
                    if any(v is not None for v in cells): table.append(dict(zip(header, cells)))
                    if len(table) > 10000: raise ValueError('Maximum 10000 rows')
            finally: workbook.close()
        else: raise HTTPException(422, 'Gunakan file CSV atau XLSX.')
        if not table or len(table) > 10000: raise ValueError('Use 1–10000 rows')
        parsed = []
        for index, row in enumerate(table, 2):
            raw = {str(k).strip(): v for k, v in row.items() if k is not None and v not in ('', None)}
            if not {'date', 'value'}.issubset(raw): raise ValueError(f'Baris {index}: date dan value wajib diisi')
            if isinstance(raw['date'], datetime): raw['date'] = raw['date'].date()
            try: parsed.append(PointInput.model_validate(raw))
            except ValueError: raise ValueError(f'Baris {index}: format tanggal/nilai/confidence tidak valid')
        return parsed
    except (ValueError, UnicodeError, KeyError, StopIteration, TypeError) as exc:
        raise HTTPException(422, 'Impor ditolak: ' + str(exc)[:250]) from exc
    except HTTPException: raise
    except Exception as exc: raise HTTPException(422, 'File tidak dapat dibaca.') from exc


@router.post('/series/{id}/import')
async def import_observations(id: str, replace_existing: bool = False, file: UploadFile = File(...), user=Depends(current_user), db=Depends(db_session)):
    require(user, 'market:write'); series = get(db, user, id, 'price', lock=True)
    if series.data.get('is_demo'): raise HTTPException(409, 'Pilih seri nyata untuk impor.')
    content = await file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024: raise HTTPException(413, 'Ukuran maksimum impor 5 MB.')
    observations = parse_import(content, file.filename or '')
    result = upsert_points(db, user, series, PointsInput(observations=observations, replace_existing=replace_existing))
    db.commit(); return result


def csv_cell(value):
    value = str(value)
    return "'" + value if value.lstrip().startswith(('=', '+', '-', '@')) else value


@router.get('/series/{id}/export')
def export_series(id: str, user=Depends(current_user), db=Depends(db_session)):
    require(user, 'market:read'); row = get(db, user, id, 'price')
    output = io.StringIO(); writer = csv.writer(output)
    writer.writerow(['date', 'value', 'confidence', 'source', 'notes'])
    for p in points(db, row): writer.writerow([p['date'], p['value'], p['confidence'], csv_cell(p['source']), csv_cell(p['notes'])])
    return Response(output.getvalue(), media_type='text/csv; charset=utf-8', headers={'Content-Disposition': 'attachment; filename="market-series.csv"'})


@router.get('/comparison')
def compare_series(ids: str, user=Depends(current_user), db=Depends(db_session)):
    require(user, 'market:read'); selected = list(dict.fromkeys(ids.split(',')))
    if not 2 <= len(selected) <= 4: raise HTTPException(422, 'Pilih 2–4 seri untuk perbandingan.')
    series = [get(db, user, id, 'price') for id in selected]
    dimensions = lambda r: (r.data.get('commodity', r.data['name']), r.data.get('currency', ''), r.data.get('unit', ''), r.data.get('metric_type', 'price'))
    if any(dimensions(row) != dimensions(series[0]) for row in series[1:]):
        raise HTTPException(422, 'Perbandingan memerlukan commodity, currency, unit, dan jenis metrik yang sama. Konversi otomatis tidak diterapkan.')
    data = [series_public(db, row, include_points=True) for row in series]
    common = set.intersection(*[{p['date'] for p in row['observations']} for row in data])
    if not common: raise HTTPException(409, 'Tidak ada tanggal observasi yang sama.')
    as_of = max(common)
    values = [next(p['value'] for p in row['observations'] if p['date'] == as_of) for row in data]
    return {'as_of': as_of, 'series': data, 'gaps': [{'series_id': row.id, 'value': value, 'vs_first_pct': change(value, values[0])} for row, value in zip(series, values)]}


class EventInput(StrictInput):
    title: str = Field(min_length=5, max_length=200)
    category_id: str
    description: str = Field(min_length=10, max_length=10000)
    date: date
    impact: Literal['Low', 'Medium', 'High'] = 'Medium'
    type: str = Field(default='Market event', min_length=3, max_length=100)
    source: str = Field(min_length=3, max_length=500)
    geography: str = Field(default='', max_length=100)
    supplier_id: str = ''
    expected_impact: str = Field(default='', max_length=3000)
    probability: float = Field(default=0.5, ge=0, le=1)
    assessment: str = Field(default='', max_length=5000)


@router.get('/events')
def list_events(category: str = '', include_demo: bool = True, user=Depends(current_user), db=Depends(db_session)):
    require(user, 'market:read'); return [public(row) for row in records(db, user, 'event', category, include_demo)]


def save_event(db, user, body, row=None):
    require(user, 'market:write'); get(db, user, body.category_id, 'category')
    if body.supplier_id: get(db, user, body.supplier_id, 'supplier')
    if row and row.data.get('is_demo'): raise HTTPException(409, 'Data contoh tidak dapat diubah.')
    row = row or Record(kind='event', owner=user.id, region=user.region)
    row.data = body.model_dump(mode='json'); row.updated_at = now()
    db.add(row); db.flush(); audit(db, user, 'market.event.saved', row.id); db.commit(); return public(row)


@router.post('/events')
def create_event(body: EventInput, user=Depends(current_user), db=Depends(db_session)):
    return save_event(db, user, body)


@router.put('/events/{id}')
def edit_event(id: str, body: EventInput, user=Depends(current_user), db=Depends(db_session)):
    return save_event(db, user, body, get(db, user, id, 'event', lock=True))


class LandscapeInput(StrictInput):
    classification: Literal['Leader', 'Challenger', 'Specialist', 'Emerging Supplier', 'New Entrant', 'Local Supplier', 'Global Supplier']
    rationale: str = Field(min_length=10, max_length=5000)
    source_ids: list[str] = Field(default_factory=list, max_length=20)


@router.put('/suppliers/{id}/landscape')
def classify_supplier(id: str, body: LandscapeInput, user=Depends(current_user), db=Depends(db_session)):
    require(user, 'market:write'); row = get(db, user, id, 'supplier', lock=True)
    for source in body.source_ids: get(db, user, source)
    if row.data.get('is_demo'): raise HTTPException(409, 'Klasifikasi hanya dapat diubah pada supplier nyata.')
    history = row.data.get('landscape_history', [])
    decision = {'classification': body.classification, 'rationale': body.rationale, 'source_ids': body.source_ids, 'approved_by': user.id, 'at': now()}
    row.data = {**row.data, 'landscape': body.classification, 'landscape_history': [*history, decision]}
    row.updated_at = now(); audit(db, user, 'market.landscape.approved', row.id, {'classification': body.classification}); db.commit()
    return public(row)


@router.get('/dashboard')
def dashboard(category: str = '', include_demo: bool = False, user=Depends(current_user), db=Depends(db_session)):
    require(user, 'market:read')
    series = [series_public(db, row) for row in records(db, user, 'price', category, include_demo)]
    events = records(db, user, 'event', category, include_demo)
    suppliers = records(db, user, 'supplier', category, include_demo)
    docs = records(db, user, 'document', category, include_demo)
    rfis = records(db, user, 'rfi', category, include_demo)
    return {'series': series, 'events': [public(row) for row in events[:12]],
            'supplier_landscape': [{'classification': label, 'count': sum(r.data.get('landscape') == label for r in suppliers)} for label in LANDSCAPES],
            'rfi_status': [{'status': status, 'count': sum(r.data.get('status') == status for r in rfis)} for status in sorted({r.data.get('status') for r in rfis})],
            'counts': {'series': len(series), 'observations': sum(s['statistics']['count'] for s in series), 'events': len(events),
                       'suppliers': len(suppliers), 'documents': len(docs), 'rfis': len(rfis),
                       'open_alerts': sum(r.data.get('status') == 'OPEN' for r in records(db, user, 'market_alert', category, include_demo))},
            'include_demo': include_demo, 'as_of': now()}
