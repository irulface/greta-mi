import hashlib, os, re, time, asyncio, zipfile
from pathlib import Path
from sqlalchemy import select, delete
from .db import Session, Record, Chunk, DATA, now, uid, Config
from .intelligence import embed
FORMATS={'pdf','docx','xlsx','pptx','txt','csv','html','htm'}
MAX_SIZE=25*1024*1024

def parse(path,ext):
    if ext in ('docx','xlsx','pptx'):
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist())>100*1024*1024:raise ValueError('Expanded document exceeds 100 MB limit.')
    if ext=='pdf':
        from pypdf import PdfReader
        reader=PdfReader(path)
        return [(i+1,p.extract_text() or '') for i,p in enumerate(reader.pages)]
    if ext=='docx':
        from docx import Document
        doc=Document(path)
        return [(1,'\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows]))]
    if ext=='xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=True)
        pages=[(i+1,s.title+'\n'+'\n'.join(' | '.join(str(v) if v is not None else '' for v in row) for row in s.iter_rows(values_only=True))) for i,s in enumerate(book.worksheets)]
        book.close();return pages
    if ext=='pptx':
        from pptx import Presentation
        return [(i+1,'\n'.join(s.text for s in slide.shapes if hasattr(s,'text'))) for i,slide in enumerate(Presentation(path).slides)]
    content=Path(path).read_text(errors='replace')
    if ext in ('html','htm'):
        from bs4 import BeautifulSoup
        soup=BeautifulSoup(content,'html.parser')
        for el in soup(['script','style']):el.decompose()
        content=soup.get_text(' ',strip=True)
    return [(1,content)]

async def index_file(db,path,name,owner,region,category='',classification='INTERNAL',source_path=None,force=False):
    content=path.read_bytes(); checksum=hashlib.sha256(content).hexdigest()
    existing=next((r for r in db.scalars(select(Record).where(Record.kind=='document',Record.region==region)) if r.data.get('checksum')==checksum and r.classification==classification and (classification!='RESTRICTED' or r.owner==owner)),None)
    if existing and not force:return 'duplicate'
    previous=existing or next((r for r in db.scalars(select(Record).where(Record.kind=='document',Record.region==region,Record.owner==owner)) if source_path and r.data.get('source_path')==source_path),None)
    ext=name.rsplit('.',1)[-1].lower();pages=parse(path,ext);text='\n'.join(x[1] for x in pages)
    target=DATA/'files'/f'{checksum}.{ext}';target.parent.mkdir(exist_ok=True)
    if not target.exists():target.write_bytes(content)
    if previous:
        db.add(Record(kind='document_version',owner=owner,region=region,classification=classification,data={'document_id':previous.id,**previous.data}))
        db.execute(delete(Chunk).where(Chunk.document_id==previous.id));doc=previous
    else:
        doc=Record(kind='document',owner=owner,region=region,classification=classification,data={});db.add(doc);db.flush()
    year=re.search(r'20\d{2}',name);rfi=re.search(r'RFI[-_ ]\d{4}[-_ ]\d{3,5}',text,re.I)
    doc.data={'name':name,'category_id':category,'format':ext.upper(),'status':'OCR REQUIRED' if not text.strip() and ext=='pdf' else 'Indexed','size':len(content),'version':doc.data.get('version',0)+1,'checksum':checksum,'storage_path':str(target),'source_path':source_path,'indexed_at':now(),'text':text[:1000000],'year':int(year.group()) if year else int(now()[:4]),'rfi_number':rfi.group() if rfi else '', 'confidence':'Requires analyst review','source':'Directory scan' if source_path else 'File upload','has_file':True,'embedding_status':'Not configured'}
    chunks=[]
    for page,value in pages:
        for start in range(0,len(value),1200):
            piece=value[max(0,start-150):start+1200]
            if piece.strip():chunks.append(Chunk(document_id=doc.id,content=piece,page=page))
    c=db.get(Config,'azure')
    if c and c.data.get('enabled') and c.data.get('embedding') and chunks:
        try:
            for start in range(0,len(chunks),16):
                batch=chunks[start:start+16];vectors=await embed(db,[c.content for c in batch])
                for chunk,vec in zip(batch,vectors):chunk.embedding=vec
            doc.data={**doc.data,'embedding_status':'Ready'}
        except Exception:doc.data={**doc.data,'embedding_status':'Failed — reindex required'}
    db.add_all(chunks);db.commit();return 'indexed'

async def run_job(job_id):
    with Session() as db:
        job=db.get(Record,job_id)
        if not job:return
        info=job.data;job.data={**info,'status':'PROCESSING','started_at':now()};db.commit()
        counts={'indexed':0,'duplicate':0,'failed':0};errors=[]
        if info['type']=='upload':files=[(Path(info['path']),info['name'])]
        else:
            root=Path(info['path']).resolve();allowed=Path(os.getenv('REPOSITORY_ROOT',str(DATA/'repository'))).resolve()
            if not root.is_relative_to(allowed):
                job.data={**job.data,'status':'FAILED','error':'Directory outside configured repository root.'};db.commit();return
            files=[(p,p.name) for p in root.rglob('*') if p.is_file() and p.resolve().is_relative_to(allowed) and not p.is_symlink() and p.suffix[1:].lower() in FORMATS]
        for path,name in files:
            try:
                if path.stat().st_size>MAX_SIZE:raise ValueError('File exceeds 25 MB limit.')
                result=await index_file(db,path,name,job.owner,job.region,info.get('category_id',''),info.get('classification','INTERNAL'),str(path) if info['type']=='scan' else None,info.get('force',False));counts[result]+=1
            except Exception as e:
                db.rollback();counts['failed']+=1;errors.append({'file':name,'error':type(e).__name__})
        job=db.get(Record,job_id);job.data={**job.data,'status':'COMPLETED' if not counts['failed'] else 'COMPLETED_WITH_ERRORS',**counts,'errors':errors,'finished_at':now()};db.commit()
        if info['type']=='upload':Path(info['path']).unlink(missing_ok=True)

async def worker():
    while True:
        try:
            with Session() as db:
                pending=db.scalar(select(Record).where(Record.kind=='job',Record.data['status'].as_string()=='PENDING').with_for_update(skip_locked=True).limit(1))
                jobs=[pending.id] if pending else []
                if pending:pending.data={**pending.data,'status':'PROCESSING','started_at':now()};db.commit()
        except Exception:
            await asyncio.sleep(3);continue
        for job in jobs:await run_job(job)
        from .mailer import process_pending_email
        try:
            await asyncio.to_thread(process_pending_email)
        except Exception:
            import logging
            logging.getLogger(__name__).exception("Email worker failed before completing a delivery")
        try:
            from .market_automation import market_tick
            await market_tick()
        except Exception:
            import logging
            logging.getLogger(__name__).exception("Market automation could not complete its cycle")
        try:
            from .enterprise import process_workflow
            await process_workflow()
        except Exception:
            import logging
            logging.getLogger(__name__).exception("Enterprise workflow cycle failed")
        await asyncio.sleep(2)
if __name__=='__main__':asyncio.run(worker())
